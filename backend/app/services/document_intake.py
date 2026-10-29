"""Safe, ordinary-paragraph DOCX extraction for the Day 2 draft-intake API."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import zipfile
import zlib

from docx import Document
from docx.opc.exceptions import PackageNotFoundError


MAX_DOCX_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_DOCX_BYTES = 20 * 1024 * 1024
MAX_DOCX_ENTRIES = 200
MAX_EXTRACTED_CHARACTERS = 100_000
ALLOWED_DOCX_EXTENSION = ".docx"


class DocumentIntakeError(ValueError):
    """A client-safe extraction error that never exposes document internals."""


class DocumentIntakeTooLargeError(DocumentIntakeError):
    """Raised when compressed, expanded, or extracted document content is too large."""


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    paragraph_count: int


def validate_docx_upload(filename: str | None, content: bytes) -> None:
    if not filename or not filename.lower().endswith(ALLOWED_DOCX_EXTENSION):
        raise DocumentIntakeError("Only .docx documents are supported.")
    if not content:
        raise DocumentIntakeError("The uploaded document is empty.")
    if len(content) > MAX_DOCX_BYTES:
        raise DocumentIntakeTooLargeError("The uploaded document exceeds the 5 MiB limit.")
    try:
        with zipfile.ZipFile(BytesIO(content)) as package:
            entries = package.infolist()
    except (zipfile.BadZipFile, zipfile.LargeZipFile, OSError, ValueError) as error:
        raise DocumentIntakeError("The upload is not a valid .docx document.") from error

    if len(entries) > MAX_DOCX_ENTRIES:
        raise DocumentIntakeTooLargeError("The document contains too many archive entries.")
    if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_DOCX_BYTES:
        raise DocumentIntakeTooLargeError("The expanded document exceeds the 20 MiB safety limit.")
def extract_docx_bytes(filename: str | None, content: bytes) -> ExtractedDocument:
    """Extract ordered paragraph text without interpreting macros or instructions."""

    validate_docx_upload(filename, content)
    try:
        document = Document(BytesIO(content))
        paragraphs = tuple(paragraph.text for paragraph in document.paragraphs)
    except (
        PackageNotFoundError,
        zipfile.BadZipFile,
        zlib.error,
        EOFError,
        OSError,
        RuntimeError,
        ValueError,
        KeyError,
        TypeError,
    ) as error:
        raise DocumentIntakeError("The .docx document could not be read.") from error

    text = "\n".join(paragraphs)
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise DocumentIntakeTooLargeError(
            "The extracted document exceeds the 100,000-character analysis limit."
        )
    if not any(paragraph.strip() for paragraph in paragraphs):
        raise DocumentIntakeError("The .docx document does not contain readable paragraph text.")
    return ExtractedDocument(text=text, paragraph_count=len(paragraphs))

